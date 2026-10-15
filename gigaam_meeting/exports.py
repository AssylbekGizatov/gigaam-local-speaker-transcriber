"""Local, explicit transcript exports. Existing files are never replaced."""
from __future__ import annotations

from datetime import datetime
from dataclasses import asdict
from html import escape
from io import BytesIO
from pathlib import Path
import json
import sys
from typing import Iterable

from .contracts import Segment, SOURCE_LABELS
from .storage import validate_local_path

EXPORT_FORMATS = ("TXT", "MD", "DOCX", "PDF", "JSON")
TIMING_NOTE = (
    "Время приблизительное, отсчитывается от начала сеанса и включает паузы. "
    "Названия источников не обозначают отдельных людей. "
    "Метки говорящих — оценка голосов, а не имена или точное число участников. "
    "Предварительные и неопределённые метки явно указаны; "
    "при перекрывающейся речи текст не разделён на независимые голоса. "
    "Текст распознавания может содержать ошибки."
)


def format_time(seconds: float) -> str:
    millis = max(0, round(seconds * 1000))
    hours, millis = divmod(millis, 3_600_000)
    minutes, millis = divmod(millis, 60_000)
    whole, millis = divmod(millis, 1000)
    return f"{hours:02d}:{minutes:02d}:{whole:02d}.{millis:03d}"


def speaker_label(segment: Segment) -> str:
    """Readable estimate separate from the capture-device label.

    Do not display a confidence score as a calibrated percentage. Unknown and
    simultaneous voices remain visible instead of receiving invented IDs.
    """
    if segment.speaker_status == "off":
        return "Диаризация выключена"
    provisional = " (предварительно)" if segment.speaker_status == "provisional" else ""
    if segment.overlapping_speakers:
        identities = " + ".join(f"Спикер {identity}" for identity in segment.overlapping_speakers)
        return f"Перекрывающаяся речь: {identities}{provisional}"
    if segment.speaker_status == "overlap":
        return "Перекрывающаяся речь: говорящие не определены"
    if segment.speaker_id is None or segment.speaker_status == "unknown":
        return f"Неопределённый говорящий{provisional}"
    return f"Спикер {segment.speaker_id}{provisional}"


def _label(segment: Segment) -> str:
    return (
        f"[{format_time(segment.start)} — {format_time(segment.end)}] "
        f"Источник: {SOURCE_LABELS.get(segment.source, str(segment.source))} | "
        f"Говорящий: {speaker_label(segment)}"
    )


def _ordered(segments: Iterable[Segment]) -> list[Segment]:
    # Stable sort: segments from both sources share one session clock.
    return sorted(segments, key=lambda segment: (segment.start, segment.end, segment.source))


def _pdf_font() -> Path:
    roots = [Path(__file__).resolve().parents[1]]
    if getattr(sys, "frozen", False):
        roots.insert(0, Path(sys.executable).resolve().parent)
    if getattr(sys, "_MEIPASS", None):
        roots.insert(0, Path(sys._MEIPASS))
    for root in roots:
        font = root / "fonts" / "DejaVuSans.ttf"
        if font.is_file():
            return font
    raise FileNotFoundError(
        "Для PDF не найден локальный шрифт fonts/DejaVuSans.ttf. "
        "Распакуйте приложение полностью. Скачивание шрифта не выполняется."
    )


def _text_bytes(segments: list[Segment], markdown: bool) -> bytes:
    title = "# Локальная транскрипция" if markdown else "Локальная транскрипция"
    lines = [title, "", TIMING_NOTE, ""]
    for segment in segments:
        label = _label(segment)
        if markdown:
            # Escape Markdown syntax from arbitrary recognized speech.
            text = segment.text.replace("\\", "\\\\")
            for char in "`*_{}[]<>#":
                text = text.replace(char, "\\" + char)
            lines.extend([f"**{label}**", "", text, ""])
        else:
            lines.extend([label, segment.text, ""])
    return ("\n".join(lines) + "\n").encode("utf-8")


def _json_bytes(segments: list[Segment]) -> bytes:
    """Lossless structured timings/identity metadata; no fabricated words.

    The four reading-oriented formats show turn intervals. JSON additionally
    retains every original Word time and distinguishes source from speaker.
    """
    payload = {
        "schemaVersion": 1,
        "timeUnit": "seconds_from_session_start_including_pauses",
        "speakerConfidenceInterpretation": "uncalibrated_score_not_probability",
        "note": TIMING_NOTE,
        "segments": [asdict(segment) for segment in segments],
    }
    return (json.dumps(payload, ensure_ascii=False, allow_nan=False, indent=2) + "\n").encode("utf-8")


def _docx_bytes(segments: list[Segment]) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    document.core_properties.title = "Локальная транскрипция"
    document.core_properties.author = ""
    document.core_properties.comments = ""
    document.add_heading("Локальная транскрипция", 0)
    document.add_paragraph(TIMING_NOTE)
    for segment in segments:
        paragraph = document.add_paragraph()
        paragraph.add_run(_label(segment)).bold = True
        document.add_paragraph(segment.text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(segments: list[Segment]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    # Embedded local Unicode font covers Kazakh, Russian and English.
    font_name = "GigaAMDejaVuSans"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(_pdf_font())))
    body = ParagraphStyle(
        "body", fontName=font_name, fontSize=10, leading=15,
        alignment=TA_LEFT, spaceAfter=8, splitLongWords=True,
    )
    heading = ParagraphStyle("heading", parent=body, fontSize=17, leading=22)
    label = ParagraphStyle("label", parent=body, textColor=colors.HexColor("#40556c"))
    note = ParagraphStyle("note", parent=body, fontSize=9, leading=13)
    flow = [Paragraph("Локальная транскрипция", heading), Spacer(1, 4 * mm)]
    flow.extend([Paragraph(escape(TIMING_NOTE), note), Spacer(1, 4 * mm)])
    for segment in segments:
        flow.append(Paragraph(escape(_label(segment)), label))
        flow.append(Paragraph(escape(segment.text).replace("\n", "<br/>"), body))
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
        title="Локальная транскрипция", author="",
    )
    document.build(flow)
    return buffer.getvalue()


def export_transcript(segments: Iterable[Segment], path: Path | str, format: str) -> Path:
    """Write an explicitly selected local destination using exclusive creation.

    DOCX/PDF are constructed in RAM. No transcript temporary files or network
    services are used. A failed write removes only the newly created file.
    """
    kind = format.upper().lstrip(".")
    if kind not in EXPORT_FORMATS:
        raise ValueError("Поддерживаются TXT, MD, DOCX, PDF и JSON.")
    destination = validate_local_path(path)
    if destination.exists():
        raise FileExistsError(f"Файл уже существует: {destination.name}. Выберите другое имя.")
    records = _ordered(segments)
    if kind == "TXT":
        payload = _text_bytes(records, markdown=False)
    elif kind == "MD":
        payload = _text_bytes(records, markdown=True)
    elif kind == "DOCX":
        payload = _docx_bytes(records)
    elif kind == "JSON":
        payload = _json_bytes(records)
    else:
        payload = _pdf_bytes(records)
    created = False
    try:
        with destination.open("xb") as stream:
            created = True
            stream.write(payload)
    except BaseException:
        if created:
            destination.unlink(missing_ok=True)
        raise
    return destination


def default_export_name(format: str) -> str:
    return "Транскрипция_" + datetime.now().strftime("%Y-%m-%d_%H-%M-%S") + "." + format.lower()
